"""Build the Idea Marketplace export workbook from data.json (extract.mjs)."""
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

DATA, OUT = sys.argv[1], sys.argv[2]
D = json.load(open(DATA))
rows, sources, bridges = D['rows'], D['sources'], D['bridges']
N = len(rows)

# ── House style ────────────────────────────────────────────────────────────
NAVY, ACCENT, TEAL, INK, MUTED = '1F3A5F', '2E86AB', '0F766E', '1F2937', '6B7280'
BAND, LINE, PANEL, WHITE = 'F4F7FB', 'D9E1EA', 'EEF3F8', 'FFFFFF'
F = 'Arial'
font = lambda **k: Font(name=F, **{'size': 10, 'color': INK, **k})
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FONT = Font(name=F, size=10, bold=True, color=WHITE)
HDR_FILL = fill(NAVY)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
wb.remove(wb.active)

# ── Column model for idea sheets ───────────────────────────────────────────
base_cols = [k for k in rows[0].keys()]
# GBP columns are FORMULAS on the stated figures and one FX input, inserted
# after the stated annual/per-vehicle figures.
gbp_after = {
    'Annual saving high (stated ccy)': ['Annual saving low (£)', 'Annual saving high (£)', 'Annual saving mid (£)'],
    'Per-vehicle high (stated ccy)': ['Per-vehicle mid (£)'],
}
cols = []
for c in base_cols:
    cols.append(c)
    cols.extend(gbp_after.get(c, []))
COL = {c: i + 1 for i, c in enumerate(cols)}
CL = {c: L(i) for c, i in COL.items()}

WIDTH = {
    'Idea ID': 16, 'Title': 46, 'Commodity': 12, 'Sub-system': 22, 'Focus area / assembly': 26, 'Level': 12,
    'Powertrain (as stated)': 16, 'Propulsion tags (derived)': 20, 'ICE': 6, 'MHEV': 7, 'PHEV': 7, 'BEV': 6, '800V': 7,
    'Off-road': 10, 'Saving type': 18, 'Annual saving (as stated)': 30, 'Stated currency': 9,
    'Annual saving low (£)': 14, 'Annual saving high (£)': 14, 'Annual saving mid (£)': 15, 'Per-vehicle mid (£)': 13,
    'Difficulty': 11, 'Time to implement': 18, 'Time min (months)': 10, 'Quick win': 9, 'Confidence': 12, 'Verified': 9,
    'Stars': 7, 'Engine check': 15, 'Engine saving %': 10,
}
LONG = {'Description', 'Technical description', 'Cost-reduction mechanism', 'Manufacturing impact', 'DFM / DFA', 'Risk notes',
        'Calculation basis', 'Volume basis', 'Engine check basis', 'Engine reference case', 'Eng: mechanism', 'Eng: spec deltas',
        'Eng: validation plan', 'Eng: DFM implications', 'Eng: cost bridge', 'Cost bridge lines', 'Review open points',
        'Evidence sources', 'Other details', 'Benchmark reference', 'Regulatory context', 'Saving rationale'}
GBP0 = '£#,##0;(£#,##0);"–"'
GBP2 = '£#,##0.00;(£#,##0.00);"–"'
EUR0 = '€#,##0;(€#,##0);"–"'
NUM_FMT = {
    'Annual saving low (stated ccy)': '#,##0', 'Annual saving high (stated ccy)': '#,##0',
    'Per-vehicle low (stated ccy)': '#,##0.00', 'Per-vehicle high (stated ccy)': '#,##0.00',
    'Annual saving low (£)': GBP0, 'Annual saving high (£)': GBP0, 'Annual saving mid (£)': GBP0, 'Per-vehicle mid (£)': GBP2,
    'Tooling (€)': EUR0, 'Capex (€)': EUR0, 'Validation (€)': EUR0, 'Payback (months)': '0.0', 'Engine saving %': '0.0%',
    'Time min (months)': '0',
}

def gbp_formula(c, r):
    cur = f'${CL["Stated currency"]}{r}'
    conv = lambda ref: f'IF({ref}="","",IF({cur}="EUR",{ref}*FX_EUR_GBP,IF({cur}="GBP",{ref},"")))'
    lo, hi = f'{CL["Annual saving low (stated ccy)"]}{r}', f'{CL["Annual saving high (stated ccy)"]}{r}'
    if c == 'Annual saving low (£)': return '=' + conv(lo)
    if c == 'Annual saving high (£)': return '=' + conv(hi)
    if c == 'Annual saving mid (£)':
        a, b = f'{CL["Annual saving low (£)"]}{r}', f'{CL["Annual saving high (£)"]}{r}'
        return f'=IF(OR({a}="",{b}=""),"",({a}+{b})/2)'
    if c == 'Per-vehicle mid (£)':
        plo, phi = f'{CL["Per-vehicle low (stated ccy)"]}{r}', f'{CL["Per-vehicle high (stated ccy)"]}{r}'
        return f'=IF({plo}="","",IF({cur}="EUR",({plo}+{phi})/2*FX_EUR_GBP,IF({cur}="GBP",({plo}+{phi})/2,"")))'

def back_link(ws, cell='A1'):
    pass

def idea_sheet(title, data, table_name, helpers=False):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = ACCENT if title != 'All Ideas' else NAVY
    allc = cols + (['Matches dashboard filters', 'Rank score'] if helpers else [])
    for j, c in enumerate(allc, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HDR_FONT; cell.fill = HDR_FILL if j <= len(cols) else fill(MUTED); cell.alignment = CENTER; cell.border = box
    for i, rw in enumerate(data, 2):
        for j, c in enumerate(cols, 1):
            if c in ('Annual saving low (£)', 'Annual saving high (£)', 'Annual saving mid (£)', 'Per-vehicle mid (£)'):
                v = gbp_formula(c, i)
            else:
                v = rw.get(c)
                if isinstance(v, str) and len(v) > 32000: v = v[:32000] + ' …[truncated at Excel cell limit]'
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = font()
            cell.alignment = WRAP_TOP
            if c in NUM_FMT: cell.number_format = NUM_FMT[c]
        if helpers:
            r = i
            prop = (f'OR(F_PROP="All",AND(F_PROP="ICE",{CL["ICE"]}{r}=1),AND(F_PROP="MHEV",{CL["MHEV"]}{r}=1),'
                    f'AND(F_PROP="PHEV",{CL["PHEV"]}{r}=1),AND(F_PROP="BEV",{CL["BEV"]}{r}=1),AND(F_PROP="800V BEV",{CL["800V"]}{r}=1),'
                    f'AND(F_PROP="Powertrain-agnostic",{CL["ICE"]}{r}+{CL["MHEV"]}{r}+{CL["PHEV"]}{r}+{CL["BEV"]}{r}=0))')
            f_ = (f'=IF(AND(OR(F_COM="All",{CL["Commodity"]}{r}=F_COM),OR(F_LVL="All",{CL["Level"]}{r}=F_LVL),'
                  f'OR(F_DIF="All",{CL["Difficulty"]}{r}=F_DIF),OR(F_OFF="All",{CL["Off-road"]}{r}=F_OFF),'
                  f'OR(F_CONF="All",{CL["Confidence"]}{r}=F_CONF),OR(F_QW="All",{CL["Quick win"]}{r}=F_QW),{prop}),1,0)')
            m = ws.cell(row=r, column=len(cols) + 1, value=f_); m.font = font(color=MUTED)
            mcol = L(len(cols) + 1)
            sc = ws.cell(row=r, column=len(cols) + 2,
                         value=f'=IF({mcol}{r}=1,N({CL["Annual saving mid (£)"]}{r})+ROW()/1000000,-1)')
            sc.font = font(color=MUTED); sc.number_format = '#,##0.000000'
    last = len(data) + 1
    ncol = len(allc)
    for c, j in COL.items():
        ws.column_dimensions[L(j)].width = WIDTH.get(c, 60 if c in LONG else 24)
    if helpers:
        ws.column_dimensions[L(len(cols) + 1)].width = 12; ws.column_dimensions[L(len(cols) + 2)].width = 14
    ws.row_dimensions[1].height = 34
    for i in range(2, last + 1): ws.row_dimensions[i].height = 48
    ws.freeze_panes = 'C2'
    tbl = Table(displayName=table_name, ref=f'A1:{L(ncol)}{last}')
    tbl.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    ws.add_table(tbl)
    # Conditional formatting: savings bars, difficulty, engine verdict, confidence.
    rng = lambda c: f'{CL[c]}2:{CL[c]}{last}'
    ws.conditional_formatting.add(rng('Annual saving mid (£)'), DataBarRule(start_type='min', end_type='percentile', end_value=95, color='7FB3D5', showValue=True))
    for val, col in (('Low', 'DCFCE7'), ('Medium', 'FEF3C7'), ('High', 'FEE2E2')):
        ws.conditional_formatting.add(rng('Difficulty'), CellIsRule(operator='equal', formula=[f'"{val}"'], fill=fill(col)))
    ws.conditional_formatting.add(rng('Engine check'), CellIsRule(operator='equal', formula=['"Confirmed"'], fill=fill('DCFCE7')))
    ws.conditional_formatting.add(rng('Engine check'), CellIsRule(operator='equal', formula=['"Contradicted"'], fill=fill('FEE2E2')))
    ws.conditional_formatting.add(rng('Quick win'), CellIsRule(operator='equal', formula=['"Yes"'], fill=fill('DCFCE7')))
    ws.conditional_formatting.add(rng('Off-road'), CellIsRule(operator='equal', formula=['"Yes"'], fill=fill('E0F2FE')))
    ws.print_title_rows = '1:1'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

COMMODITIES = ['Battery', 'EDU', 'Electrical', 'Powertrain', 'Driveline', 'Chassis', 'BIW', 'Exterior', 'Interior', 'Other']
LEVELS = ['Assembly', 'Subassembly', 'Part', 'System', 'Not stated']
DIFFS = ['Low', 'Medium', 'High']
CONFS = ['verified', 'benchmarked', 'estimated', 'Not stated']
PROPS = ['ICE', 'MHEV', 'PHEV', 'BEV', '800V BEV', 'Powertrain-agnostic']
assert set(r['Commodity'] for r in rows) <= set(COMMODITIES), set(r['Commodity'] for r in rows)
assert set(r['Level'] for r in rows) <= set(LEVELS), set(r['Level'] for r in rows)
assert set(r['Confidence'] for r in rows) <= set(CONFS), set(r['Confidence'] for r in rows)

# ── Sheets in order ────────────────────────────────────────────────────────
cover = wb.create_sheet('Cover')
dash = wb.create_sheet('Dashboard')
allws = idea_sheet('All Ideas', rows, 'tblAllIdeas', helpers=True)
import os
LITE = os.environ.get('LITE')
for c in COMMODITIES:
    sub = [r for r in rows if r['Commodity'] == c]
    if LITE: sub = sub[:2]
    idea_sheet(c, sub, f'tbl{c.replace(" ", "")}')

# Sources
ws = wb.create_sheet('Sources'); ws.sheet_properties.tabColor = TEAL
scols = ['Idea ID', 'Idea title', 'Commodity', 'Source title', 'Type', 'Year', 'Confidence', 'URL', 'Supports']
for j, c in enumerate(scols, 1):
    x = ws.cell(row=1, column=j, value=c); x.font = HDR_FONT; x.fill = HDR_FILL; x.alignment = CENTER
for i, s_ in enumerate(sources, 2):
    for j, c in enumerate(scols, 1):
        x = ws.cell(row=i, column=j, value=s_[c]); x.font = font(); x.alignment = WRAP_TOP
        if c == 'URL' and isinstance(s_[c], str) and s_[c].startswith('http'):
            x.hyperlink = s_[c]; x.font = font(color='1D4ED8', underline='single')
for j, w in enumerate([16, 40, 12, 40, 14, 7, 11, 50, 70], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'C2'
t = Table(displayName='tblSources', ref=f'A1:{L(len(scols))}{len(sources) + 1}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); ws.add_table(t)

# Cost bridges
ws = wb.create_sheet('Cost Bridges'); ws.sheet_properties.tabColor = TEAL
bcols = ['Idea ID', 'Idea title', 'Commodity', 'Item', 'Baseline (€)', 'Proposed (€)', 'Saving (€)', 'Saving (£)', 'Note']
for j, c in enumerate(bcols, 1):
    x = ws.cell(row=1, column=j, value=c); x.font = HDR_FONT; x.fill = HDR_FILL; x.alignment = CENTER
for i, b in enumerate(bridges, 2):
    for j, c in enumerate(bcols, 1):
        v = f'=G{i}*FX_EUR_GBP' if c == 'Saving (£)' else (f'=E{i}-F{i}' if c == 'Saving (€)' else b.get(c))
        x = ws.cell(row=i, column=j, value=v); x.font = font(); x.alignment = WRAP_TOP
        if c in ('Baseline (€)', 'Proposed (€)', 'Saving (€)'): x.number_format = '€#,##0.00;(€#,##0.00);"–"'
        if c == 'Saving (£)': x.number_format = '£#,##0.00;(£#,##0.00);"–"'
for j, w in enumerate([16, 40, 12, 40, 13, 13, 13, 13, 50], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'C2'
t = Table(displayName='tblCostBridges', ref=f'A1:{L(len(bcols))}{len(bridges) + 1}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); ws.add_table(t)

# Lists (helper; hidden) — dropdown sources and the filtered sub-system ranking
lists = wb.create_sheet('Lists')
def put_list(col, title, vals):
    lists.cell(row=1, column=col, value=title).font = font(bold=True)
    for i, v in enumerate(['All'] + vals, 2): lists.cell(row=i, column=col, value=v)
    return f"Lists!${L(col)}$2:${L(col)}${len(vals) + 2}"
LIST_RNG = {
    'F_COM': put_list(1, 'Commodity', COMMODITIES), 'F_PROP': put_list(2, 'Propulsion', PROPS),
    'F_LVL': put_list(3, 'Level', LEVELS), 'F_DIF': put_list(4, 'Difficulty', DIFFS),
    'F_OFF': put_list(5, 'Off-road', ['Yes', 'No', 'Not stated']), 'F_CONF': put_list(6, 'Confidence', CONFS),
    'F_QW': put_list(7, 'Quick win', ['Yes', 'No']),
}
last = N + 1
A = lambda c: f"'All Ideas'!${CL[c]}$2:${CL[c]}${last}"
MATCH = f"'All Ideas'!${L(len(cols) + 1)}$2:${L(len(cols) + 1)}${last}"
SCORE = f"'All Ideas'!${L(len(cols) + 2)}$2:${L(len(cols) + 2)}${last}"
MID = A('Annual saving mid (£)')
systems = sorted(set(r['Sub-system'] for r in rows))
lists.cell(row=1, column=9, value='Sub-system').font = font(bold=True)
lists.cell(row=1, column=10, value='Ideas (filtered)').font = font(bold=True)
lists.cell(row=1, column=11, value='Rank key').font = font(bold=True)
for i, sname in enumerate(systems, 2):
    lists.cell(row=i, column=9, value=sname)
    lists.cell(row=i, column=10, value=f'=COUNTIFS({A("Sub-system")},I{i},{MATCH},1)')
    lists.cell(row=i, column=11, value=f'=J{i}+ROW()/100000')
SYS_LAST = len(systems) + 1
lists.sheet_state = 'hidden'

# ── Dashboard ──────────────────────────────────────────────────────────────
d = dash
d.sheet_properties.tabColor = 'C9A227'
d.sheet_view.showGridLines = False
for c, w in zip('ABCDEFGHIJKLMNOPQRSTU', [2, 26, 14, 14, 14, 14, 12, 26, 14, 14, 14, 14, 3, 34, 14, 14, 14, 14, 14, 14, 3]):
    d.column_dimensions[c].width = w
def band(rng, color=PANEL):
    for row in d[rng]:
        for c in row: c.fill = fill(color)
d.merge_cells('B1:S1')
d['B1'] = 'Idea Marketplace — Cost-Reduction Opportunity Dashboard'
d['B1'].font = Font(name=F, size=18, bold=True, color=WHITE); d['B1'].fill = fill(NAVY); d['B1'].alignment = Alignment(vertical='center', indent=1)
d.row_dimensions[1].height = 36
d.merge_cells('B2:S2')
d['B2'] = ('Every figure is a live formula on the All Ideas tab — change a filter and the whole page recalculates. '
           'Savings are engineering estimates per idea, in £; they are NOT additive (many ideas are alternatives), so totals are a gross sum of independent estimates at each idea\'s own stated volume, not a programme saving.')
d['B2'].font = font(size=9, italic=True, color=MUTED); d['B2'].alignment = Alignment(wrap_text=True, vertical='top', indent=1)
d.row_dimensions[2].height = 30

# Filters panel
d['B4'] = 'FILTERS'; d['B4'].font = font(bold=True, color=NAVY, size=11)
filters = [('Commodity', 'F_COM'), ('Propulsion', 'F_PROP'), ('Level', 'F_LVL'), ('Difficulty', 'F_DIF'),
           ('Off-road', 'F_OFF'), ('Confidence', 'F_CONF'), ('Quick wins', 'F_QW')]
for i, (lab, nm) in enumerate(filters):
    r = 5 + i
    d.cell(row=r, column=2, value=lab).font = font(bold=True)
    c = d.cell(row=r, column=3, value='All'); c.font = Font(name=F, size=10, bold=True, color='0000FF'); c.fill = fill('FFF9DB'); c.border = box; c.alignment = Alignment(horizontal='center')
    d.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    dv = DataValidation(type='list', formula1=LIST_RNG[nm], allow_blank=False, showDropDown=False)
    dv.error = 'Pick a value from the list'; dv.prompt = f'Filter the dashboard by {lab.lower()}'; d.add_data_validation(dv); dv.add(c)
    wb.defined_names[nm] = DefinedName(nm, attr_text=f'Dashboard!$C${r}')
d['B12'] = 'Yellow cells are inputs: pick from the dropdown; "All" clears a filter.'; d['B12'].font = font(size=8, italic=True, color=MUTED)
d['B14'] = 'ASSUMPTION'; d['B14'].font = font(bold=True, color=NAVY, size=11)
d['B15'] = '£ per €1 (EUR→GBP)'; d['B15'].font = font(bold=True)
fx = d['C15']; fx.value = 0.85; fx.number_format = '0.0000'; fx.font = Font(name=F, size=10, bold=True, color='0000FF'); fx.fill = fill('FFF9DB'); fx.border = box
fx.comment = Comment('Source: BrainSpark fx-rates.mjs reference table (FX_FALLBACK GBP 0.85 per EUR, as of 1 Sep 2026). No live ECB rate was reachable when this workbook was built. Change this cell and every £ figure recalculates.', 'BrainSpark')
wb.defined_names['FX_EUR_GBP'] = DefinedName('FX_EUR_GBP', attr_text='Dashboard!$C$15')
d['B16'] = 'Converts the ideas stated in €; ideas stated in £ are used as stated.'; d['B16'].font = font(size=8, italic=True, color=MUTED)

# KPI tiles (row 4–7, columns E..S in pairs)
kpis = [
    ('Ideas in view', f'=SUM({MATCH})', '#,##0'),
    ('Gross annual savings (£, not additive)', f'=SUMIFS({MID},{MATCH},1)', '£#,##0,,"M"'),
    ('Median idea (£/yr)', f'=IFERROR(MEDIAN(IF({MATCH}=1,IF(ISNUMBER({MID}),{MID}))),0)', '£#,##0'),
    ('Quick wins', f'=COUNTIFS({MATCH},1,{A("Quick win")},"Yes")', '#,##0'),
    ('Engine-confirmed', f'=COUNTIFS({MATCH},1,{A("Engine check")},"Confirmed")', '#,##0'),
    ('Off-road ideas', f'=COUNTIFS({MATCH},1,{A("Off-road")},"Yes")', '#,##0'),
]
tile_cols = [(6, 7), (8, 9), (10, 11), (14, 15), (16, 17), (18, 19)]
for (lab, f_, fmt), (c1, c2) in zip(kpis, tile_cols):
    d.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    d.merge_cells(start_row=5, start_column=c1, end_row=7, end_column=c2)
    h = d.cell(row=4, column=c1, value=lab); h.font = font(size=8, bold=True, color=MUTED); h.alignment = CENTER
    v = d.cell(row=5, column=c1, value=f_); v.font = Font(name=F, size=20, bold=True, color=NAVY); v.alignment = CENTER; v.number_format = fmt
    for rr in range(4, 8):
        for cc in range(c1, c2 + 1):
            d.cell(row=rr, column=cc).fill = fill(PANEL); d.cell(row=rr, column=cc).border = Border(top=thin if rr == 4 else None, bottom=thin if rr == 7 else None, left=thin if cc == c1 else None, right=thin if cc == c2 else None)
# Median is an array formula
from openpyxl.worksheet.formula import ArrayFormula
mc = d.cell(row=5, column=10)
mc.value = ArrayFormula(f'{L(10)}5', f'=IFERROR(MEDIAN(IF({MATCH}=1,IF(ISNUMBER({MID}),{MID}))),0)')
d['F8'] = 'Share of all ideas in view'; d['F8'].font = font(size=8, color=MUTED)
d['H8'] = f'=IFERROR(F5/{N},0)'; d['H8'].number_format = '0.0%'; d['H8'].font = font(size=8, bold=True, color=MUTED)

def section(r, c, title):
    x = d.cell(row=r, column=c, value=title); x.font = font(bold=True, color=NAVY, size=11)

def header(r, c, labels):
    for k, lab in enumerate(labels):
        x = d.cell(row=r, column=c + k, value=lab); x.font = HDR_FONT; x.fill = HDR_FILL; x.alignment = CENTER; x.border = box

def grid_row(r, c, vals, fmts):
    for k, (v, f_) in enumerate(zip(vals, fmts)):
        x = d.cell(row=r, column=c + k, value=v); x.font = font(); x.border = box
        if f_: x.number_format = f_
        if k == 0: x.font = font(bold=True)

# By commodity (B19)
R0 = 19
section(R0 - 1, 2, 'By commodity')
header(R0, 2, ['Commodity', 'Ideas', 'Savings (£M)', 'Quick wins', 'Share of view'])
for i, c in enumerate(COMMODITIES):
    r = R0 + 1 + i
    grid_row(r, 2, [c, f'=COUNTIFS({A("Commodity")},B{r},{MATCH},1)', f'=SUMIFS({MID},{A("Commodity")},B{r},{MATCH},1)/1000000',
                    f'=COUNTIFS({A("Commodity")},B{r},{MATCH},1,{A("Quick win")},"Yes")', f'=IFERROR(C{r}/$F$5,0)'],
             [None, '#,##0', '£#,##0.0', '#,##0', '0.0%'])
COM_LAST = R0 + len(COMMODITIES)
r = COM_LAST + 1
grid_row(r, 2, ['Total', f'=SUM(C{R0+1}:C{COM_LAST})', f'=SUM(D{R0+1}:D{COM_LAST})', f'=SUM(E{R0+1}:E{COM_LAST})', f'=IFERROR(C{r}/$F$5,0)'], [None, '#,##0', '£#,##0.0', '#,##0', '0.0%'])
for k in range(5): d.cell(row=r, column=2 + k).font = font(bold=True); d.cell(row=r, column=2 + k).fill = fill(PANEL)
d.conditional_formatting.add(f'C{R0+1}:C{COM_LAST}', DataBarRule(start_type='num', start_value=0, end_type='max', color='7FB3D5', showValue=True))
d.conditional_formatting.add(f'D{R0+1}:D{COM_LAST}', DataBarRule(start_type='num', start_value=0, end_type='max', color='7DCEA0', showValue=True))

# By propulsion (H19)
section(R0 - 1, 8, 'By propulsion (an idea can serve several)')
header(R0, 8, ['Propulsion', 'Ideas', 'Savings (£M)', 'Quick wins', 'Share of view'])
flagcol = {'ICE': 'ICE', 'MHEV': 'MHEV', 'PHEV': 'PHEV', 'BEV': 'BEV', '800V BEV': '800V'}
for i, p in enumerate(PROPS):
    r = R0 + 1 + i
    if p in flagcol:
        cnt = f'=COUNTIFS({A(flagcol[p])},1,{MATCH},1)'; sm = f'=SUMIFS({MID},{A(flagcol[p])},1,{MATCH},1)/1000000'
        qw = f'=COUNTIFS({A(flagcol[p])},1,{MATCH},1,{A("Quick win")},"Yes")'
    else:
        cnt = f'=COUNTIFS({A("ICE")},0,{A("MHEV")},0,{A("PHEV")},0,{A("BEV")},0,{MATCH},1)'
        sm = f'=SUMIFS({MID},{A("ICE")},0,{A("MHEV")},0,{A("PHEV")},0,{A("BEV")},0,{MATCH},1)/1000000'
        qw = f'=COUNTIFS({A("ICE")},0,{A("MHEV")},0,{A("PHEV")},0,{A("BEV")},0,{MATCH},1,{A("Quick win")},"Yes")'
    grid_row(r, 8, [p, cnt, sm, qw, f'=IFERROR(I{r}/$F$5,0)'], [None, '#,##0', '£#,##0.0', '#,##0', '0.0%'])
PROP_LAST = R0 + len(PROPS)
d.conditional_formatting.add(f'I{R0+1}:I{PROP_LAST}', DataBarRule(start_type='num', start_value=0, end_type='max', color='7FB3D5', showValue=True))

# By level (H27)
L0 = PROP_LAST + 3
section(L0 - 1, 8, 'By level (assembly / subassembly / part)')
header(L0, 8, ['Level', 'Ideas', 'Savings (£M)', 'Quick wins', 'Share of view'])
for i, lv in enumerate(LEVELS):
    r = L0 + 1 + i
    grid_row(r, 8, [lv, f'=COUNTIFS({A("Level")},H{r},{MATCH},1)', f'=SUMIFS({MID},{A("Level")},H{r},{MATCH},1)/1000000',
                    f'=COUNTIFS({A("Level")},H{r},{MATCH},1,{A("Quick win")},"Yes")', f'=IFERROR(I{r}/$F$5,0)'],
             [None, '#,##0', '£#,##0.0', '#,##0', '0.0%'])
LEV_LAST = L0 + len(LEVELS)

# Difficulty + confidence + engine (N19)
section(R0 - 1, 14, 'By difficulty')
header(R0, 14, ['Difficulty', 'Ideas', 'Savings (£M)', 'Share of view'])
for i, dv_ in enumerate(DIFFS):
    r = R0 + 1 + i
    grid_row(r, 14, [dv_, f'=COUNTIFS({A("Difficulty")},N{r},{MATCH},1)', f'=SUMIFS({MID},{A("Difficulty")},N{r},{MATCH},1)/1000000', f'=IFERROR(O{r}/$F$5,0)'], [None, '#,##0', '£#,##0.0', '0.0%'])
DIF_LAST = R0 + len(DIFFS)
C0 = DIF_LAST + 3
section(C0 - 1, 14, 'By evidence confidence')
header(C0, 14, ['Confidence', 'Ideas', 'Savings (£M)', 'Share of view'])
for i, cv in enumerate(CONFS):
    r = C0 + 1 + i
    grid_row(r, 14, [cv, f'=COUNTIFS({A("Confidence")},N{r},{MATCH},1)', f'=SUMIFS({MID},{A("Confidence")},N{r},{MATCH},1)/1000000', f'=IFERROR(O{r}/$F$5,0)'], [None, '#,##0', '£#,##0.0', '0.0%'])
CONF_LAST = C0 + len(CONFS)
E0 = CONF_LAST + 3
section(E0 - 1, 14, 'Engine cross-check')
header(E0, 14, ['Engine check', 'Ideas', 'Savings (£M)', 'Share of view'])
for i, ev in enumerate(['Confirmed', 'Contradicted', 'Not engine-checked']):
    r = E0 + 1 + i
    grid_row(r, 14, [ev, f'=COUNTIFS({A("Engine check")},N{r},{MATCH},1)', f'=SUMIFS({MID},{A("Engine check")},N{r},{MATCH},1)/1000000', f'=IFERROR(O{r}/$F$5,0)'], [None, '#,##0', '£#,##0.0', '0.0%'])
ENG_LAST = E0 + 3

# Commodity × level matrix
M0 = max(COM_LAST, LEV_LAST, ENG_LAST) + 4
section(M0 - 1, 2, 'Commodity × level — number of ideas (heat map)')
header(M0, 2, ['Commodity'] + LEVELS + ['Total'])
for i, c in enumerate(COMMODITIES):
    r = M0 + 1 + i
    vals = [c] + [f'=COUNTIFS({A("Commodity")},$B{r},{A("Level")},{L(3 + k)}${M0},{MATCH},1)' for k in range(len(LEVELS))] + [f'=SUM(C{r}:{L(2 + len(LEVELS))}{r})']
    grid_row(r, 2, vals, [None] + ['#,##0'] * (len(LEVELS) + 1))
MAT_LAST = M0 + len(COMMODITIES)
d.conditional_formatting.add(f'C{M0+1}:{L(2+len(LEVELS))}{MAT_LAST}', ColorScaleRule(start_type='min', start_color='FFFFFF', mid_type='percentile', mid_value=50, mid_color='BFD7EA', end_type='max', end_color='2E86AB'))

# Top 15 sub-systems (filtered) — columns N..P
S0 = M0
section(S0 - 1, 14, 'Top 15 sub-systems in view')
header(S0, 14, ['Sub-system', 'Ideas', 'Savings (£M)'])
for k in range(1, 16):
    r = S0 + k
    key = f'LARGE(Lists!$K$2:$K${SYS_LAST},{k})'
    name = f'INDEX(Lists!$I$2:$I${SYS_LAST},MATCH({key},Lists!$K$2:$K${SYS_LAST},0))'
    grid_row(r, 14, [f'=IF(INT({key})=0,"",{name})', f'=IF(N{r}="","",COUNTIFS({A("Sub-system")},N{r},{MATCH},1))',
                     f'=IF(N{r}="","",SUMIFS({MID},{A("Sub-system")},N{r},{MATCH},1)/1000000)'], [None, '#,##0', '£#,##0.0'])
    d.cell(row=r, column=14).font = font()
d.conditional_formatting.add(f'O{S0+1}:O{S0+15}', DataBarRule(start_type='num', start_value=0, end_type='max', color='7FB3D5', showValue=True))
SUB_LAST = S0 + 15

# Top 20 ideas (filtered) — own tab, same filters
tp = wb.create_sheet('Top 20'); tp.sheet_properties.tabColor = 'C9A227'; tp.sheet_view.showGridLines = False
tp.merge_cells('B1:M1'); tp['B1'] = 'Top 20 ideas in view — ranked by annual saving (£, mid of the stated range)'
tp['B1'].font = Font(name=F, size=16, bold=True, color=WHITE); tp['B1'].fill = fill(NAVY); tp['B1'].alignment = Alignment(vertical='center', indent=1)
tp.row_dimensions[1].height = 32
tp['B2'] = '=CONCATENATE("Filters from the Dashboard — Commodity: ",F_COM,"  ·  Propulsion: ",F_PROP,"  ·  Level: ",F_LVL,"  ·  Difficulty: ",F_DIF,"  ·  Off-road: ",F_OFF,"  ·  Confidence: ",F_CONF,"  ·  Quick wins: ",F_QW)'
tp['B2'].font = font(size=9, italic=True, color=MUTED)
tp['B3'] = 'Back to Dashboard'; tp['B3'].hyperlink = "#'Dashboard'!A1"; tp['B3'].font = font(color='1D4ED8', underline='single', size=9)
tcols = ['#', 'Idea ID', 'Title', 'Commodity', 'Sub-system', 'Level', 'Propulsion', 'Annual saving (£)', 'Difficulty', 'Time to implement', 'Engine check', 'Confidence']
tsrc = [None, 'Idea ID', 'Title', 'Commodity', 'Sub-system', 'Level', 'Propulsion tags (derived)', 'Annual saving mid (£)', 'Difficulty', 'Time to implement', 'Engine check', 'Confidence']
T0 = 5
for k_, lab in enumerate(tcols):
    x = tp.cell(row=T0, column=2 + k_, value=lab); x.font = HDR_FONT; x.fill = HDR_FILL; x.alignment = CENTER; x.border = box
for k in range(1, 21):
    r = T0 + k
    key = f'LARGE({SCORE},{k})'
    rowno = f'MATCH({key},{SCORE},0)'
    vals = [k] + [f'=IF({key}<0,"",INDEX({A(src)},{rowno}))' for src in tsrc[1:]]
    for k_, v in enumerate(vals):
        x = tp.cell(row=r, column=2 + k_, value=v); x.font = font(bold=(k_ == 2)); x.border = box; x.alignment = WRAP_TOP
        if k % 2 == 0: x.fill = fill(BAND)
    tp.cell(row=r, column=9).number_format = '£#,##0'
    tp.row_dimensions[r].height = 32
for c_, w in zip('ABCDEFGHIJKLM', [2, 5, 16, 60, 12, 24, 12, 22, 15, 11, 20, 15, 12]):
    tp.column_dimensions[c_].width = w
tp.conditional_formatting.add(f'I{T0+1}:I{T0+20}', DataBarRule(start_type='num', start_value=0, end_type='max', color='7DCEA0', showValue=True))
for val, col in (('Low', 'DCFCE7'), ('Medium', 'FEF3C7'), ('High', 'FEE2E2')):
    tp.conditional_formatting.add(f'J{T0+1}:J{T0+20}', CellIsRule(operator='equal', formula=[f'"{val}"'], fill=fill(col)))
tp.conditional_formatting.add(f'L{T0+1}:L{T0+20}', CellIsRule(operator='equal', formula=['"Confirmed"'], fill=fill('DCFCE7')))
tp.conditional_formatting.add(f'L{T0+1}:L{T0+20}', CellIsRule(operator='equal', formula=['"Contradicted"'], fill=fill('FEE2E2')))
tp.page_setup.orientation = 'landscape'; tp.page_setup.fitToWidth = 1; tp.sheet_properties.pageSetUpPr.fitToPage = True
d['B13'] = 'Top 20 ideas for these filters →'; d['B13'].hyperlink = "#'Top 20'!A1"; d['B13'].font = font(color='1D4ED8', underline='single', bold=True)
TOP_LAST = max(MAT_LAST, SUB_LAST)

# Charts
def bar(title, cats, vals, anchor, w=17, h=8.5, horizontal=True, color=ACCENT, fmt='#,##0'):
    ch = BarChart(); ch.type = 'bar' if horizontal else 'col'; ch.style = 10
    ch.title = title; ch.legend = None; ch.height = h; ch.width = w
    ch.add_data(vals, titles_from_data=False); ch.set_categories(cats)
    s_ = ch.series[0]; s_.graphicalProperties.solidFill = color; s_.graphicalProperties.line.solidFill = color
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False; ch.dataLabels.numFmt = fmt
    ch.y_axis.majorGridlines = None; ch.y_axis.numFmt = fmt
    if horizontal: ch.x_axis.scaling.orientation = 'maxMin'
    ch.gapWidth = 60
    ch.x_axis.tickLblSkip = 1; ch.x_axis.delete = False; ch.y_axis.delete = False
    d.add_chart(ch, anchor)
CH0 = TOP_LAST + 3
section(CH0 - 1, 2, 'Charts (follow the filters)')
bar('Ideas by commodity', Reference(d, min_col=2, min_row=R0 + 1, max_row=COM_LAST), Reference(d, min_col=3, min_row=R0 + 1, max_row=COM_LAST), f'B{CH0}')
bar('Gross annual savings by commodity (£M, not additive)', Reference(d, min_col=2, min_row=R0 + 1, max_row=COM_LAST), Reference(d, min_col=4, min_row=R0 + 1, max_row=COM_LAST), f'H{CH0}', color=TEAL, fmt='£#,##0')
bar('Ideas by propulsion', Reference(d, min_col=8, min_row=R0 + 1, max_row=PROP_LAST), Reference(d, min_col=9, min_row=R0 + 1, max_row=PROP_LAST), f'N{CH0}', w=15)
CH1 = CH0 + 19
bar('Ideas by level', Reference(d, min_col=8, min_row=L0 + 1, max_row=LEV_LAST), Reference(d, min_col=9, min_row=L0 + 1, max_row=LEV_LAST), f'B{CH1}', horizontal=False, color=NAVY)
bar('Top 15 sub-systems in view (ideas)', Reference(d, min_col=14, min_row=S0 + 1, max_row=SUB_LAST), Reference(d, min_col=15, min_row=S0 + 1, max_row=SUB_LAST), f'H{CH1}', h=10)
pie = PieChart(); pie.title = 'Ideas by difficulty'; pie.height = 8.5; pie.width = 15
pie.add_data(Reference(d, min_col=15, min_row=R0 + 1, max_row=DIF_LAST), titles_from_data=False)
pie.set_categories(Reference(d, min_col=14, min_row=R0 + 1, max_row=DIF_LAST))
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True; pie.dataLabels.showVal = False; pie.dataLabels.showSerName = False; pie.dataLabels.showCatName = True; pie.dataLabels.showLegendKey = False
from openpyxl.chart.series import DataPoint
for idx, col in enumerate(['22C55E', 'F59E0B', 'EF4444']):
    pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = col; pie.series[0].dPt.append(pt)
d.add_chart(pie, f'N{CH1}')
d.freeze_panes = 'A4'
d.print_area = f'A1:S{CH1 + 24}'
d.page_setup.orientation = 'landscape'; d.page_setup.fitToWidth = 1; d.page_setup.fitToHeight = 0
d.sheet_properties.pageSetUpPr.fitToPage = True

# ── Cover ──────────────────────────────────────────────────────────────────
cv = cover; cv.sheet_view.showGridLines = False; cv.sheet_properties.tabColor = NAVY
cv.column_dimensions['A'].width = 3; cv.column_dimensions['B'].width = 34; cv.column_dimensions['C'].width = 18; cv.column_dimensions['D'].width = 90
for r in range(1, 6):
    for c in range(1, 6): cv.cell(row=r, column=c).fill = fill(NAVY)
cv.column_dimensions['E'].width = 3
cv['B2'] = 'BrainSpark Idea Marketplace'; cv['B2'].font = Font(name=F, size=24, bold=True, color=WHITE)
cv['B3'] = 'Cost-reduction idea library — full export with dashboard'; cv['B3'].font = Font(name=F, size=13, color='D6E4F0')
cv['B4'] = f'Prepared {datetime.date(2026, 10, 12).strftime("%d %B %Y")} · for the Cost Engineering Director and President review'; cv['B4'].font = Font(name=F, size=10, color='D6E4F0')
cv['B7'] = 'Scope'; cv['B7'].font = font(bold=True, color=NAVY, size=12)
facts = [
    ('Ideas exported', f"=COUNTA('All Ideas'!A2:A{last})", 'Every live idea in the marketplace library (14 packs); ideas retired after engineering review are excluded.'),
    ('Fields per idea', len(cols), 'Every stored field, plus clearly labelled derived columns (commodity, propulsion tags, £ figures, quick win).'),
    ('Evidence sources', f"=COUNTA(Sources!A2:A{len(sources)+1})", 'One row per web or industry source behind a benchmark fact (Sources tab, clickable links).'),
    ('Cost-bridge lines', f"=COUNTA('Cost Bridges'!A2:A{len(bridges)+1})", 'Baseline → proposed € per line for ideas that carry a bridge (Cost Bridges tab).'),
    ('Currency', '£ (GBP)', 'Engines compute in €; € figures are converted at the rate on the Dashboard (editable). Ideas stated in £ are used as stated.'),
]
for i, (a, b, c) in enumerate(facts, 8):
    cv.cell(row=i, column=2, value=a).font = font(bold=True)
    x = cv.cell(row=i, column=3, value=b); x.font = Font(name=F, size=11, bold=True, color=NAVY); x.alignment = Alignment(horizontal='left'); x.number_format = '#,##0'
    cv.cell(row=i, column=4, value=c).font = font(color=MUTED); cv.cell(row=i, column=4).alignment = Alignment(wrap_text=True)
cv['B14'] = 'Contents'; cv['B14'].font = font(bold=True, color=NAVY, size=12)
toc = [('Dashboard', 'Interactive summary: filters, KPIs, breakdowns, top sub-systems, charts'), ('Top 20', 'The 20 largest ideas for the Dashboard filters'),
       ('All Ideas', 'Master table — every idea, every field; filterable Excel table')] + \
      [(c, f'{c} ideas — same columns as All Ideas') for c in COMMODITIES] + \
      [('Sources', 'Evidence sources per idea, with links'), ('Cost Bridges', 'Per-line cost bridges'), ('Field Guide', 'What every column means and where it comes from')]
for i, (sh, desc) in enumerate(toc, 15):
    x = cv.cell(row=i, column=2, value=sh); x.hyperlink = f"#'{sh}'!A1"; x.font = font(color='1D4ED8', underline='single', bold=True)
    if sh in COMMODITIES:
        cv.cell(row=i, column=3, value=f"=COUNTA('{sh}'!A:A)-1").number_format = '#,##0 "ideas"'
        cv.cell(row=i, column=3).font = font(color=MUTED)
    cv.cell(row=i, column=4, value=desc).font = font(color=MUTED)
NR = 15 + len(toc) + 1
cv.cell(row=NR, column=2, value='How to read the numbers').font = font(bold=True, color=NAVY, size=12)
notes = [
    'Savings are ENGINEERING ESTIMATES per idea at the volume each idea states (most at 50,000–100,000 vehicles/yr). They are not quotes and not commitments.',
    'Savings are NOT ADDITIVE: many ideas are alternatives to each other or assume different architectures. Dashboard totals are a gross sum of independent estimates, not a programme saving.',
    '"Annual saving (as stated)" is the idea\'s own text; the £ columns are parsed from it (low / high / mid). An idea whose saving has no figure shows "–".',
    '"Engine check" is the deterministic should-cost engine re-pricing the move on a reference case: Confirmed = the engine agrees the direction saves money; it does not verify this part\'s exact figure.',
    '"Sources flagged unreviewed" = Yes means the benchmark sources were found by AI web search and have not been reviewed by an engineer; they support the benchmark fact, never the saving.',
    'Propulsion tags are derived from each idea\'s stated powertrain and its text (an idea can carry several). "Powertrain-agnostic" = none detected.',
    '"Review open points" are findings from the engineering review that need an engineer\'s judgement before the idea is used.',
]
for i, n in enumerate(notes, NR + 1):
    x = cv.cell(row=i, column=2, value=f'•  {n}'); x.font = font(); x.alignment = Alignment(wrap_text=True, vertical='top')
    cv.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4); cv.row_dimensions[i].height = 34
cv.page_setup.orientation = 'portrait'; cv.page_setup.fitToWidth = 1; cv.page_setup.fitToHeight = 1
cv.sheet_properties.pageSetUpPr.fitToPage = True
cv.print_area = f'A1:E{NR + len(notes)}'

# ── Field guide ────────────────────────────────────────────────────────────
fg = wb.create_sheet('Field Guide'); fg.sheet_properties.tabColor = MUTED
GUIDE = {
    'Idea ID': 'Unique marketplace id.', 'Title': 'Idea title.',
    'Commodity': 'DERIVED — the app\'s commodity classifier applied to the sub-system (same mapping as the Marketplace commodity tabs).',
    'Sub-system': 'The system the idea attacks (marketplace "system").', 'Focus area / assembly': 'The assembly or focus area named by the idea.',
    'Level': 'Assembly / Subassembly / Part / System as stored; "Not stated" for legacy ideas that carry no level.',
    'Powertrain (as stated)': 'Powertrain the idea was written for, where stated.',
    'Propulsion tags (derived)': 'DERIVED — powertrains and 800V detected from the stated powertrain and the idea text (the marketplace facet classifier).',
    'ICE': 'DERIVED flag (1 = applies).', 'MHEV': 'DERIVED flag.', 'PHEV': 'DERIVED flag.', 'BEV': 'DERIVED flag.', '800V': 'DERIVED flag.',
    'Off-road': 'Off-road capability idea (Yes/No; Not stated for packs that do not carry the flag).',
    'Annual saving (as stated)': 'The idea\'s own annual-saving text, verbatim.', 'Stated currency': 'Currency of that text (EUR or GBP).',
    'Annual saving low (stated ccy)': 'DERIVED — low end parsed from the stated text, in the stated currency.',
    'Annual saving high (stated ccy)': 'DERIVED — high end parsed from the stated text.',
    'Annual saving low (£)': 'FORMULA — low in £ (× the Dashboard EUR→GBP rate when stated in €).',
    'Annual saving high (£)': 'FORMULA — high in £.', 'Annual saving mid (£)': 'FORMULA — midpoint of low and high in £; drives the dashboard.',
    'Per-vehicle mid (£)': 'FORMULA — midpoint of the stated per-vehicle saving in £, where stated.',
    'Time min (months)': 'DERIVED — first number in the time-to-implement text.',
    'Quick win': 'DERIVED — Yes when difficulty is Low and implementation can start within 12 months.',
    'Engine check': 'Deterministic should-cost engine cross-check: Confirmed / Contradicted / Not engine-checked.',
    'Engine saving %': 'The engine\'s saving on its reference case (not this part\'s figure).',
    'Other details': 'Any stored field not given its own column (kept so nothing is lost; empty for almost every idea).',
    'Library pack': 'Which seeded idea pack the idea came from.',
    'Matches dashboard filters': 'HELPER — 1 when the row passes every Dashboard filter.', 'Rank score': 'HELPER — annual £ mid for ranking the Top 20 (−1 when filtered out).',
}
for j, c in enumerate(['Column', 'Meaning / source'], 1):
    x = fg.cell(row=1, column=j, value=c); x.font = HDR_FONT; x.fill = HDR_FILL
for i, c in enumerate(cols + ['Matches dashboard filters', 'Rank score'], 2):
    fg.cell(row=i, column=1, value=c).font = font(bold=True)
    x = fg.cell(row=i, column=2, value=GUIDE.get(c, 'Stored idea field, exported verbatim.')); x.font = font(); x.alignment = Alignment(wrap_text=True)
fg.column_dimensions['A'].width = 34; fg.column_dimensions['B'].width = 110; fg.freeze_panes = 'A2'

# Order: Cover, Dashboard, All Ideas, commodities..., Sources, Cost Bridges, Field Guide, Lists
order = ['Cover', 'Dashboard', 'Top 20', 'All Ideas'] + COMMODITIES + ['Sources', 'Cost Bridges', 'Field Guide', 'Lists']
wb._sheets = [wb[s] for s in order]
wb.active = 0
for wsx in wb.worksheets:
    wsx.sheet_view.zoomScale = 90 if wsx.title not in ('Dashboard', 'Cover') else 100
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, 'cols', len(cols))
