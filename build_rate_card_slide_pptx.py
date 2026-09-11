#!/usr/bin/env python3
"""
One slide for the IT team: how JLR's own rate card gets into the tool.

Separate from the five-slide board deck on purpose — the audience is different.
The director wants the case; IT wants to see the actual screen, who is allowed
to press the buttons, where the file goes and how the tool proves it is using
the new numbers instead of the built-in ones.

EVERYTHING ON THIS SLIDE WAS RUN, NOT DESCRIBED. Captured 11 September 2026
against the real server with AIR_GAPPED=1:

  admin gate    POST /api/rate-library/upload as an ordinary user -> HTTP 403
                the same call as an admin                          -> HTTP 200
  upload        328 materials, 178 machines, 42 labour, 11 energy,
                9 FX, 23 overhead defaults; version 93d84eb673273903
  before        built-in book 1398b5ef6f6758d6 -> £33.19
  after         company book  93d84eb673273903 -> £70.60
  same part, same answers, same rules version (v6) both times

The screenshot is the real Rate Library panel, not a mock-up.

Regenerate:  python3 build_rate_card_slide_pptx.py
Output:      CostVision-Rate-Card-Upload.pptx
"""

import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from pptx_fixup import finalise

# Same palette as the business-case deck so the two read as one pack.
BLUE    = RGBColor(0x1D, 0x6F, 0xB8)
DARK    = RGBColor(0x16, 0x32, 0x5C)
BODY    = RGBColor(0x3A, 0x43, 0x56)
MUTED   = RGBColor(0x6B, 0x72, 0x80)
BG      = RGBColor(0xF4, 0xF7, 0xFB)
PANEL   = RGBColor(0xFF, 0xFF, 0xFF)
GREENBG = RGBColor(0xEA, 0xF6, 0xEF)
PANEL2  = RGBColor(0xE8, 0xF1, 0xFA)
GREEN   = RGBColor(0x2E, 0x8B, 0x57)
AMBER   = RGBColor(0xB7, 0x79, 0x1F)
LINE    = RGBColor(0xDC, 0xE3, 0xEE)
NAVY    = RGBColor(0x16, 0x32, 0x5C)
ON_DARK = RGBColor(0xFF, 0xFF, 0xFF)

TITLE_FONT = 'Cambria'
SHOT = 'assets/rate-library-admin.png'

W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None, round_=False, radius=0.12):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if round_:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


_PICTO = re.compile('([\U0001F300-\U0001FAFF☀-⛿✀-✒✙-➿️]+)')


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=3, line_spacing=1.0, font='Calibri'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for r in para:
            t, size, color, bold = r[0], r[1], r[2], r[3]
            italic = r[4] if len(r) > 4 else False
            for part in _PICTO.split(t):
                if not part: continue
                run = p.add_run(); run.text = part
                f = run.font
                f.size = Pt(size); f.color.rgb = color; f.bold = bold; f.italic = italic
                f.name = 'Segoe UI Emoji' if _PICTO.fullmatch(part) else font
    return tb


slide = prs.slides.add_slide(BLANK)
box(slide, 0, 0, W, H, fill=BG)

# ── Header ───────────────────────────────────────────────────────────────────
badge = box(slide, Inches(0.35), Inches(0.22), Inches(0.42), Inches(0.42),
            fill=BLUE, round_=True, radius=0.28)
tf = badge.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'cv'
r.font.size = Pt(17); r.font.bold = True; r.font.name = 'Calibri'; r.font.color.rgb = ON_DARK
text(slide, Inches(0.87), Inches(0.19), Inches(2.6), Inches(0.32),
     [[('CostVision', 18, BLUE, True)]])
box(slide, 0, Inches(0.78), W, Pt(2.2), fill=NAVY)

text(slide, Inches(0.45), Inches(0.95), Inches(11.5), Inches(0.3),
     [[('RATE LIBRARY · ADMIN', 11, BLUE, True)]])
text(slide, Inches(0.45), Inches(1.22), Inches(12.4), Inches(0.5),
     [[('Loading our own rate card', 27, DARK, True)]], font=TITLE_FONT)
text(slide, Inches(0.45), Inches(1.72), Inches(12.4), Inches(0.28),
     [[('An admin uploads one Excel file. From that moment every costing uses our rates '
        'instead of the ones that ship with the tool.', 12.5, BODY, False)]])

# ── The flow ─────────────────────────────────────────────────────────────────
FLOW_Y = Inches(2.16)
FLOW_H = Inches(1.06)
steps = [
    ('1', 'Sign in as admin', 'Only an admin sees the\nbuttons. Anyone else gets\nrefused by the server.'),
    ('2', 'Download template', 'An Excel file with every\nrate the tool uses, filled\nwith the current values.'),
    ('3', 'Put our rates in it', 'Material £/kg, labour £/hr,\nmachine build-ups, energy,\nFX, overhead.'),
    ('4', 'Upload it', 'Checked as it loads. If a\ncolumn is wrong it is\nrejected, not half-loaded.'),
    ('5', 'Switch it on', 'One click. The tool now\nreads our book. One more\nputs it back.'),
]
gap = Inches(0.16)
bw = (Inches(12.44) - gap * (len(steps) - 1)) / len(steps)
for i, (n, head, sub) in enumerate(steps):
    x = Inches(0.45) + (bw + gap) * i
    box(slide, x, FLOW_Y, bw, FLOW_H, fill=PANEL, line=LINE)
    box(slide, x, FLOW_Y, bw, Pt(3), fill=BLUE)
    text(slide, x + Inches(0.12), FLOW_Y + Inches(0.13), Inches(0.3), Inches(0.2),
         [[(n, 11, BLUE, True)]])
    text(slide, x + Inches(0.38), FLOW_Y + Inches(0.12), bw - Inches(0.5), Inches(0.22),
         [[(head, 11.5, DARK, True)]])
    text(slide, x + Inches(0.12), FLOW_Y + Inches(0.40), bw - Inches(0.24), Inches(0.58),
         [[(l, 9.2, BODY, False)] for l in sub.split('\n')], space_after=0, line_spacing=1.06)
    if i < len(steps) - 1:
        a = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                   x + bw + Inches(0.015), FLOW_Y + Inches(0.44),
                                   Inches(0.13), Inches(0.18))
        a.rotation = 90
        a.fill.solid(); a.fill.fore_color.rgb = LINE
        a.line.fill.background(); a.shadow.inherit = False

# ── Screenshot ───────────────────────────────────────────────────────────────
SHOT_Y = Inches(3.44)
text(slide, Inches(0.45), SHOT_Y, Inches(7.6), Inches(0.24),
     [[('THE ACTUAL SCREEN', 10, MUTED, True)]])
try:
    # Height is pinned, not left to the aspect ratio: a taller crop silently grew
    # the picture down over the strip below, and a picture does not clip.
    slide.shapes.add_picture(SHOT, Inches(0.45), SHOT_Y + Inches(0.26),
                             width=Inches(7.62), height=Inches(2.21))
except Exception:
    box(slide, Inches(0.45), SHOT_Y + Inches(0.26), Inches(7.62), Inches(2.3), fill=PANEL, line=LINE)
    text(slide, Inches(0.7), SHOT_Y + Inches(1.3), Inches(7.1), Inches(0.3),
         [[(f'screenshot missing: {SHOT}', 11, MUTED, False)]])

# ── Proof panel ──────────────────────────────────────────────────────────────
PX, PW = Inches(8.32), Inches(4.57)
box(slide, PX, SHOT_Y, PW, Inches(2.62), fill=PANEL, line=LINE)
box(slide, PX, SHOT_Y, Pt(3), Inches(2.62), fill=GREEN)
text(slide, PX + Inches(0.18), SHOT_Y + Inches(0.14), PW - Inches(0.36), Inches(0.24),
     [[('Does it really use our numbers?', 12.5, GREEN, True)]])
text(slide, PX + Inches(0.18), SHOT_Y + Inches(0.42), PW - Inches(0.36), Inches(0.34),
     [[('The same bracket, the same answers, costed twice — before the upload '
        'and after it.', 10.2, BODY, False)]])

rows = [
    ('', 'Built-in', 'Our sheet', True),
    ('Material', '£12.79', '£28.56', False),
    ('Labour', '£5.36', '£16.07', False),
    ('Machine time', '£8.91', '£13.36', False),
    ('Total', '£33.19', '£70.60', True),
]
ry = SHOT_Y + Inches(0.86)
for label, a, b, bold in rows:
    if label == 'Total':
        box(slide, PX + Inches(0.18), ry - Inches(0.03), PW - Inches(0.36), Inches(0.29), fill=GREENBG)
    text(slide, PX + Inches(0.24), ry, Inches(1.7), Inches(0.22),
         [[(label, 10.4, DARK if bold else BODY, bold)]])
    text(slide, PX + Inches(2.0), ry, Inches(1.0), Inches(0.22),
         [[(a, 10.4, MUTED if not bold else DARK, bold)]], align=PP_ALIGN.RIGHT)
    text(slide, PX + Inches(3.1), ry, Inches(1.2), Inches(0.22),
         [[(b, 10.4, DARK, bold)]], align=PP_ALIGN.RIGHT)
    ry += Inches(0.28)

text(slide, PX + Inches(0.18), ry + Inches(0.04), PW - Inches(0.36), Inches(0.42),
     [[('Only the buckets a rate feeds moved. Packaging, logistics and tooling did '
        'not — nothing in the sheet touches them.', 9.4, MUTED, False, True)]],
     line_spacing=1.08)

# ── Audit strip ──────────────────────────────────────────────────────────────
AY = Inches(6.22)
box(slide, PX, AY, PW, Inches(0.84), fill=PANEL2)
text(slide, PX + Inches(0.18), AY + Inches(0.12), PW - Inches(0.36), Inches(0.22),
     [[('Every costing records which book it used', 11.5, BLUE, True)]])
text(slide, PX + Inches(0.18), AY + Inches(0.38), PW - Inches(0.36), Inches(0.4),
     [[('Each sheet gets an ID from its own contents — ours is 93d84eb6, the '
        'built-in one 1398b5ef. A report months old can be traced back to the exact '
        'rates behind it.', 9.4, BODY, False)]], line_spacing=1.08)

# ── Footnote under the screenshot ────────────────────────────────────────────
FY = Inches(6.22)
box(slide, Inches(0.45), FY, Inches(7.62), Inches(0.84), fill=PANEL, line=LINE)
box(slide, Inches(0.45), FY, Pt(3), Inches(0.84), fill=AMBER)
text(slide, Inches(0.68), FY + Inches(0.12), Inches(7.2), Inches(0.22),
     [[('Who can do this, and what happens if nobody does', 11.5, AMBER, True)]])
text(slide, Inches(0.68), FY + Inches(0.38), Inches(7.2), Inches(0.4),
     [[('Admins are named in the server settings, and the check is done on the server, '
        'not in the browser. With no sheet uploaded the tool keeps using its own '
        'rates and says so on every run.', 9.4, BODY, False)]], line_spacing=1.08)

# ── Speaker notes ────────────────────────────────────────────────────────────
slide.notes_slide.notes_text_frame.text = (
    "This is the slide for the IT team, and the short version is that it is one Excel "
    "file and two clicks. I will walk the five boxes along the top, and then show you "
    "that it actually does what it says.\n\n"
    "First, who. The buttons on that screen only appear for an admin, and more "
    "importantly the check is done on the server, not in the browser — so hiding a "
    "button is not the security. We tried it: an ordinary signed-in user calling the "
    "upload directly gets refused, and the same call from an admin account goes "
    "through. Admins are named in the server settings, so you decide who they are, and "
    "the role is read fresh every time rather than being baked into the login token. "
    "Take someone's admin off and it takes effect on their next click.\n\n"
    "Second, the file. You press Download Excel template and you get a workbook with "
    "every rate the tool uses, already filled in with what it currently has — so you "
    "are editing a real sheet, not filling in a blank one. Material prices per kilo, "
    "labour rates per hour, the machine build-ups, energy, exchange rates and the "
    "overhead defaults. You put our numbers in the same columns and upload it. It is "
    "checked as it loads — if a column is missing or a number is not a number, the "
    "whole file is rejected rather than half-loaded, because half a rate book is worse "
    "than none.\n\n"
    "The upload we did loaded three hundred and twenty-eight materials, a hundred and "
    "seventy-eight machines and forty-two labour grades. Then one click switches the "
    "tool over to it, and one click puts it back.\n\n"
    "Now the part I would want to see if I were you. The panel on the right is the same "
    "bracket costed twice, with the same answers, changing nothing except which rate "
    "book is switched on. Thirty-three pounds nineteen on the built-in rates, seventy "
    "pounds sixty on ours. Material more than doubles, labour triples, machine time "
    "goes up. And look at what did not move — packaging, logistics and tooling are "
    "unchanged, because nothing in the sheet touches them. That is the tool reading our "
    "file rather than a number being nudged somewhere.\n\n"
    "Last thing, and it matters for audit. Each sheet gets an ID worked out from its own "
    "contents, and every costing records which one it used. So a report from months ago "
    "can be traced back to the exact rates behind it, and if someone asks why two "
    "costings of the same part differ, the answer is in the record rather than in "
    "somebody's memory.\n\n"
    "If nobody ever uploads anything, nothing breaks — the tool keeps using its own "
    "rates and says so on every run. It is never guessing about which numbers it is on."
)

OUT = 'CostVision-Rate-Card-Upload.pptx'
prs.save(OUT)
finalise(OUT)
print(f'{OUT}  -  1 slide')
