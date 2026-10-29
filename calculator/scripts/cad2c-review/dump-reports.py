#!/usr/bin/env python3
"""Dump every exported report in a run directory to text for a line-by-line audit (CAD2C five-part review, Oct 2026).

  python3 -I scripts/cad2c-review/dump-reports.py <run dir> [<out dir>]

For each <label>.pdf: <label>.pdf.txt (page by page, PyMuPDF text) and <label>.p<N>.png (page renders at 110 dpi).
For each <label>.xlsx: <label>.xlsx.txt (sheet by sheet, every cell).
"""
import sys, os, glob
import pymupdf  # type: ignore
import openpyxl  # type: ignore

run = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(run, 'audit')
os.makedirs(out, exist_ok=True)
for pdf in sorted(glob.glob(os.path.join(run, '*.pdf'))):
    label = os.path.basename(pdf)[:-4]
    doc = pymupdf.open(pdf)
    with open(os.path.join(out, label + '.pdf.txt'), 'w') as f:
        for i, page in enumerate(doc):
            f.write(f'\n===== PAGE {i + 1} / {len(doc)} ===== images: {len(page.get_images())}\n')
            f.write(page.get_text('text'))
            pix = page.get_pixmap(dpi=110)
            pix.save(os.path.join(out, f'{label}.p{i + 1:02d}.png'))
    print(f'{label}.pdf: {len(doc)} pages')
for xl in sorted(glob.glob(os.path.join(run, '*.xlsx'))):
    label = os.path.basename(xl)[:-5]
    wb = openpyxl.load_workbook(xl, data_only=False)
    with open(os.path.join(out, label + '.xlsx.txt'), 'w') as f:
        for ws in wb.worksheets:
            f.write(f'[images on {ws.title}: {len(getattr(ws, "_images", []))}]\n')
            f.write(f'\n===== SHEET {ws.title} ({ws.max_row} rows × {ws.max_column} cols) =====\n')
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    f.write('\n'); continue
                f.write(' | '.join('' if v is None else str(v) for v in row).rstrip(' |') + '\n')
    print(f'{label}.xlsx: {len(wb.worksheets)} sheets')
